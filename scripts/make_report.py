
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "reports" / "REPORT_TEMPLATE.md"
OUTPUT = ROOT / "reports" / "MLOps_Assignment01_Report.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10)
for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
    styles[name].font.name = "Calibri"

lines = SOURCE.read_text(encoding="utf-8-sig").splitlines()
i = 0

def add_table(table_lines):
    rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(cols):
            cells[ci].text = row[ci] if ci < len(row) else ""
            if ri == 0:
                for run in cells[ci].paragraphs[0].runs:
                    run.bold = True

def clean(text):
    text = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    return text.strip()

while i < len(lines):
    line = lines[i].strip()

    if not line or line in ("---", "***"):
        i += 1
        continue

    if line.startswith("|") and "|" in line[1:]:
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i].strip())
            i += 1
        add_table(block)
        continue

    heading = re.match(r"^(#{1,6})\s*(.*)$", line)
    if heading:
        level = min(len(heading.group(1)), 3)
        text = clean(heading.group(2))
        if text:
            if not doc.paragraphs and level == 1:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(text)
            else:
                doc.add_heading(text, level=level)
        i += 1
        continue

    if line.startswith("```"):
        code_lines = []
        i += 1
        while i < len(lines) and not lines[i].strip().startswith("```"):
            code_lines.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run("\n".join(code_lines))
        r.font.name = "Consolas"
        r.font.size = Pt(8)
        i += 1
        continue

    if re.match(r"^[-*+]\s+", line):
        text = re.sub(r"^[-*+]\s+", "", line)
        if text.startswith("[x]") or text.startswith("[ ]"):
            text = ("☒ " if text.startswith("[x]") else "☐ ") + text[3:].strip()
        doc.add_paragraph(clean(text), style="List Bullet")
        i += 1
        continue

    if re.match(r"^\d+[.)]\s+", line):
        text = re.sub(r"^\d+[.)]\s+", "", line)
        doc.add_paragraph(clean(text), style="List Number")
        i += 1
        continue

    if line.startswith(">"):
        doc.add_paragraph(clean(line.lstrip("> ")), style="Quote")
        i += 1
        continue

    paragraph = [line]
    i += 1
    while i < len(lines):
        nxt = lines[i].strip()
        if (
            not nxt or nxt.startswith("#") or nxt.startswith("|")
            or nxt.startswith("```") or nxt in ("---", "***")
            or re.match(r"^[-*+]\s+", nxt)
            or re.match(r"^\d+[.)]\s+", nxt)
        ):
            break
        paragraph.append(nxt)
        i += 1

    doc.add_paragraph(clean(" ".join(paragraph)))

doc.core_properties.title = "MLOps Assignment 01 — Heart Disease Prediction"
doc.core_properties.subject = "MLOps Assignment Report"
doc.save(OUTPUT)
print(f"Report created: {OUTPUT}")
print(f"Source lines processed: {len(lines)}")
